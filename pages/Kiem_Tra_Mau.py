import io
import os
import sys
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
import pandas as pd
import streamlit as st

# -------------------------------------------------------------
# CẤU HÌNH GIAO DIỆN CHUẨN APPLE
# -------------------------------------------------------------
st.set_page_config(
    page_title="Đối Chiếu Nghiệm Thu HRC - KSNB",
    page_icon="🛡️",
    layout="wide",
)

st.markdown(
    """
<style>
    [data-testid="stSidebar"] { display: none !important; }
    [data-testid="collapsedControl"] { display: none !important; }
    .stApp {
        background-color: #f5f5f7;
        font-family: -apple-system, BlinkMacSystemFont, "SF Pro Display", sans-serif;
        color: #1d1d1f;
    }
    .apple-header-box {
        display: flex;
        justify-content: space-between;
        align-items: flex-end;
        padding: 20px 0 10px 0;
        margin-bottom: 5px;
    }
    .apple-main-title {
        font-size: 36px;
        font-weight: 700;
        color: #1d1d1f;
        letter-spacing: -0.025em;
        margin: 0;
    }
    div.stButton > button:first-child {
        background: #0071e3 !important;
        color: #ffffff !important;
        border-radius: 980px !important;
        border: none !important;
        padding: 10px 20px !important;
        font-weight: 600 !important;
        font-size: 14px !important;
        box-shadow: 0 4px 12px rgba(0, 113, 227, 0.2) !important;
    }
    [data-testid="stVerticalBlockBorderWrapper"] {
        background: #ffffff;
        border-radius: 16px;
        padding: 18px;
        border: 1px solid rgba(0, 0, 0, 0.08);
        box-shadow: 0 2px 12px rgba(0, 0, 0, 0.03);
    }
    .upload-title {
        font-size: 17px;
        font-weight: 700;
        color: #1d1d1f;
        margin-bottom: 14px;
    }
    .guide-box {
        background: #fbfbfd;
        border-radius: 10px;
        padding: 12px;
        margin-bottom: 12px;
        border: 1px solid rgba(0,0,0,0.04);
    }
    .guide-header {
        font-weight: 700;
        color: #0071e3;
        font-size: 13px;
        margin-bottom: 4px;
    }
    .guide-text {
        font-size: 12px;
        color: #515154;
        line-height: 1.5;
        margin: 0;
    }
</style>
""",
    unsafe_allow_html=True,
)

# Header
st.markdown(
    """
<div class="apple-header-box">
    <h1 class="apple-main-title">📑 Đối Chiếu Nghiệm Thu HRC</h1>
</div>
<hr style='border: none; border-top: 1px solid #e2e8f0; margin: 10px 0 25px 0;'>
""",
    unsafe_allow_html=True,
)


def doi_chieu_hrc_tung_file():
    st.caption("Kiểm soát nội bộ • Tự động đối soát dữ liệu và tô đỏ sai lệch")

    col_upload1, col_upload2 = st.columns(2, gap="large")

    with col_upload1:
        with st.container(border=True):
            st.markdown(
                """
            <div class="upload-title">1️⃣ Dữ liệu Nghiệm thu & Trạm cân</div>
            <div class="guide-box"><div class="guide-header">📌 1. BBNT</div><div class="guide-text">Biên bản nghiệm thu hàng thép cuộn cán nóng</div></div>
            <div class="guide-box"><div class="guide-header">📌 2. Packing List</div><div class="guide-text">Báo cáo chi tiết nhập hàng NCC</div></div>
            <div class="guide-box"><div class="guide-header">📌 3. Dữ liệu Trạm cân</div><div class="guide-text">Dữ liệu cân xe thực tế</div></div>
            """,
                unsafe_allow_html=True,
            )
            f_bbnt = st.file_uploader(
                "1. File BBNT", type=["xlsx", "xls"], key="bbnt"
            )
            f_pl = st.file_uploader(
                "2. File Packing List", type=["xlsx", "xls"], key="pl"
            )
            f_tc = st.file_uploader(
                "3. File Trạm cân", type=["xlsx", "xls"], key="tc"
            )

    with col_upload2:
        with st.container(border=True):
            st.markdown(
                """
            <div class="upload-title">2️⃣ Dữ liệu Bảo vệ & ERP</div>
            <div class="guide-box"><div class="guide-header">📌 4. Nhập Bảo vệ</div><div class="guide-text">Tra cứu xuất nhập hàng cổng bảo vệ</div></div>
            <div class="guide-box"><div class="guide-header">📌 5. Báo cáo ERP OPM 105</div><div class="guide-text">Báo cáo kiểm soát thông tin giao dịch theo lô</div></div>
            """,
                unsafe_allow_html=True,
            )
            f_bv = st.file_uploader(
                "4. File Nhập Bảo vệ", type=["xlsx", "xls"], key="bv"
            )
            f_opm = st.file_uploader(
                "5. File ERP OPM 105", type=["xlsx", "xls"], key="opm"
            )

    if f_bbnt is not None:
        if st.button("🚀 Bắt đầu đối chiếu tự động", use_container_width=True):
            with st.spinner("Đang tính toán và lập bảng đối chiếu..."):
                try:
                    # 1. ĐỌC BBNT (Header ở dòng 8 -> skiprows=7)
                    df_bbnt = pd.read_excel(f_bbnt, skiprows=7)
                    df_bbnt = df_bbnt.dropna(
                        subset=[df_bbnt.columns[3]]
                    ).copy()
                    # Loại bỏ các dòng tổng cộng / ghi chú kết luận ở cuối
                    df_bbnt = df_bbnt[
                        ~df_bbnt[df_bbnt.columns[0]]
                        .astype(str)
                        .str.contains("Tổng cộng|Kết luận", case=False, na=False)
                    ]

                    # Giữ đúng 13 cột thông tin ban đầu của BBNT
                    base_cols = df_bbnt.columns[:13].tolist()
                    res_df = df_bbnt[base_cols].copy()

                    col_ma_cuon = df_bbnt.columns[3]
                    col_phieu_can = df_bbnt.columns[4]
                    col_net_ct = df_bbnt.columns[7]
                    col_gross_ct = df_bbnt.columns[8]
                    col_can_ban = df_bbnt.columns[9]
                    col_lech_pt_bbnt = df_bbnt.columns[11]

                    ma_cuon_clean = (
                        df_bbnt[col_ma_cuon]
                        .astype(str)
                        .str.replace(r"\.0$", "", regex=True)
                        .str.strip()
                        .str.upper()
                    )

                    phieu_can_clean = (
                        df_bbnt[col_phieu_can]
                        .astype(str)
                        .str.replace(r"\.0$", "", regex=True)
                        .str.strip()
                    )

                    ma_hoa = ma_cuon_clean + "A"

                    # 2. PACKING LIST (Dòng header thực tế là dòng 8 -> skiprows=7)
                    dict_pl = {}
                    if f_pl is not None:
                        df_pl = pd.read_excel(f_pl, skiprows=7)
                        # Lấy cột theo vị trí hoặc tên: Mã số NSX (col 7), KL(net) (col 11)
                        c_pl_ma = df_pl.columns[7]
                        c_pl_net = df_pl.columns[11]

                        df_pl_clean = df_pl.dropna(
                            subset=[c_pl_ma, c_pl_net]
                        ).copy()
                        pl_key = (
                            df_pl_clean[c_pl_ma]
                            .astype(str)
                            .str.replace(r"\.0$", "", regex=True)
                            .str.strip()
                            .str.upper()
                            + "A"
                        )
                        pl_net_val = pd.to_numeric(
                            df_pl_clean[c_pl_net], errors="coerce"
                        ).fillna(0)
                        dict_pl = (
                            pd.DataFrame({"Key": pl_key, "Net": pl_net_val})
                            .groupby("Key")["Net"]
                            .sum()
                            .to_dict()
                        )

                    # 3. NHẬP BẢO VỆ
                    dict_bv_phieu = {}
                    dict_bv_kl = {}
                    if f_bv is not None:
                        df_bv = pd.read_excel(f_bv)
                        col_hang = next(
                            (
                                c
                                for c in df_bv.columns
                                if "TÊN HÀNG" in str(c).upper()
                            ),
                            None,
                        )
                        col_pcan = next(
                            (
                                c
                                for c in df_bv.columns
                                if "PHIẾU CÂN" in str(c).upper()
                                and "KHỐI" not in str(c).upper()
                            ),
                            None,
                        )
                        col_klcan = next(
                            (
                                c
                                for c in df_bv.columns
                                if "KHỐI LƯỢNG" in str(c).upper()
                            ),
                            None,
                        )

                        if col_hang and col_pcan and col_klcan:
                            for _, r in df_bv.iterrows():
                                txt = str(r[col_hang]).upper()
                                p_can = (
                                    str(r[col_pcan])
                                    .replace(".0", "")
                                    .strip()
                                )
                                kl_can = r[col_klcan]
                                for m in ma_cuon_clean:
                                    if m in txt and m not in dict_bv_phieu:
                                        dict_bv_phieu[m] = p_can
                                        dict_bv_kl[m] = kl_can

                    # 4. TRẠM CÂN
                    dict_tc_kl = {}
                    dict_tc_phanloai = {}
                    if f_tc is not None:
                        df_tc = pd.read_excel(f_tc)
                        col_tc_p = next(
                            (
                                c
                                for c in df_tc.columns
                                if str(c).strip().upper() == "PHIẾU CÂN"
                            ),
                            df_tc.columns[0],
                        )
                        col_tc_k = next(
                            (
                                c
                                for c in df_tc.columns
                                if "KHỐI LƯỢNG" in str(c).upper()
                            ),
                            None,
                        )
                        col_tc_pl = next(
                            (
                                c
                                for c in df_tc.columns
                                if "PHÂN LOẠI" in str(c).upper()
                            ),
                            None,
                        )

                        tc_p = (
                            df_tc[col_tc_p]
                            .astype(str)
                            .str.replace(r"\.0$", "", regex=True)
                            .str.strip()
                        )
                        tc_kl = (
                            pd.to_numeric(
                                df_tc[col_tc_k], errors="coerce"
                            ).fillna(0)
                            if col_tc_k
                            else 0
                        )
                        tc_pl = (
                            df_tc[col_tc_pl].fillna("Nhập hàng").astype(str)
                            if col_tc_pl
                            else "Nhập hàng"
                        )

                        dict_tc_kl = dict(zip(tc_p, tc_kl))
                        dict_tc_phanloai = dict(zip(tc_p, tc_pl))

                    # 5. OPM105 (Header thực tế là dòng 10 -> skiprows=9, dòng tiếp theo là số thứ tự [1],[2])
                    opm_035_po, opm_035_rtv, opm_031, opm_03g = {}, {}, {}, {}
                    if f_opm is not None:
                        # Dòng 10 là tiêu đề chữ
                        df_opm = pd.read_excel(f_opm, skiprows=9)
                        # Bỏ dòng chứa số thứ tự [1], [2], ...
                        if len(df_opm) > 0 and "[" in str(df_opm.iloc[0, 1]):
                            df_opm = df_opm.iloc[1:].copy()

                        # Tự động map tên cột chuẩn
                        col_opm_loai = next(
                            (
                                c
                                for c in df_opm.columns
                                if "LOẠI GIAO DỊCH" in str(c).upper()
                            ),
                            df_opm.columns[1],
                        )
                        col_opm_malo = next(
                            (
                                c
                                for c in df_opm.columns
                                if str(c).strip().upper() == "MÃ LÔ"
                            ),
                            df_opm.columns[6],
                        )
                        col_opm_kl = next(
                            (
                                c
                                for c in df_opm.columns
                                if "SỐ LƯỢNG" in str(c).upper()
                                or "KHỐI LƯỢNG" in str(c).upper()
                            ),
                            df_opm.columns[10],
                        )
                        col_opm_kho = next(
                            (
                                c
                                for c in df_opm.columns
                                if "KHO" in str(c).upper()
                            ),
                            df_opm.columns[35],
                        )

                        df_opm_clean = df_opm.dropna(
                            subset=[col_opm_malo, col_opm_loai]
                        ).copy()
                        # Chuẩn hóa mã lô: nếu chưa có chữ A thì thêm A
                        lo_clean = (
                            df_opm_clean[col_opm_malo]
                            .astype(str)
                            .str.strip()
                            .str.upper()
                        )
                        df_opm_clean["Mã Lô Key"] = lo_clean.apply(
                            lambda x: x if x.endswith("A") else x + "A"
                        )
                        df_opm_clean["KL"] = pd.to_numeric(
                            df_opm_clean[col_opm_kl], errors="coerce"
                        ).fillna(0)
                        df_opm_clean["Kho"] = (
                            df_opm_clean[col_opm_kho].astype(str).str.strip()
                        )

                        opm_035_po = (
                            df_opm_clean[
                                (df_opm_clean["Kho"] == "035")
                                & (
                                    df_opm_clean[col_opm_loai].str.contains(
                                        "PO Receipt", case=False, na=False
                                    )
                                )
                            ]
                            .groupby("Mã Lô Key")["KL"]
                            .sum()
                            .to_dict()
                        )
                        opm_035_rtv = (
                            df_opm_clean[
                                (df_opm_clean["Kho"] == "035")
                                & (
                                    df_opm_clean[col_opm_loai].str.contains(
                                        "Return to Vendor",
                                        case=False,
                                        na=False,
                                    )
                                )
                            ]
                            .groupby("Mã Lô Key")["KL"]
                            .sum()
                            .to_dict()
                        )
                        opm_031 = (
                            df_opm_clean[
                                (df_opm_clean["Kho"].isin(["03G", "03H"]))
                                & (
                                    df_opm_clean[col_opm_loai].str.contains(
                                        "Intransit Receipt",
                                        case=False,
                                        na=False,
                                    )
                                )
                            ]
                            .groupby("Mã Lô Key")["KL"]
                            .sum()
                            .to_dict()
                        )
                        opm_03g = (
                            df_opm_clean[
                                (df_opm_clean["Kho"] == "03G")
                                & (
                                    df_opm_clean[col_opm_loai].str.contains(
                                        "PO Receipt", case=False, na=False
                                    )
                                )
                            ]
                            .groupby("Mã Lô Key")["KL"]
                            .sum()
                            .to_dict()
                        )

                    # 6. TÍNH TOÁN CÁC CỘT ĐỐI SOÁT
                    res_df["NHẬP/XUẤT"] = phieu_can_clean.map(
                        lambda x: dict_tc_phanloai.get(x, "Nhập hàng")
                    )
                    res_df["Mã hóa"] = ma_hoa
                    res_df["Khối lượng net theo PL"] = (
                        ma_hoa.map(dict_pl).fillna(0).astype(float)
                    )

                    net_ct_val = (
                        pd.to_numeric(df_bbnt[col_net_ct], errors="coerce")
                        .fillna(0)
                        .astype(float)
                    )
                    res_df["Đối chiếu"] = (
                        net_ct_val - res_df["Khối lượng net theo PL"]
                    )

                    res_df["KL theo bc nhập hàng"] = (
                        ma_cuon_clean.map(dict_bv_kl).fillna(0).astype(float)
                    )
                    res_df["Số phiếu cân"] = ma_cuon_clean.map(
                        dict_bv_phieu
                    ).fillna("")
                    res_df["Đối chiếu số phiếu cân"] = (
                        phieu_can_clean == res_df["Số phiếu cân"].astype(str)
                    )

                    res_df["KL theo bc trạm cân"] = (
                        phieu_can_clean.map(dict_tc_kl).fillna(0).astype(float)
                    )
                    res_df["Đối chiếu KL TC & KL BCBV"] = (
                        res_df["KL theo bc trạm cân"]
                        == res_df["KL theo bc nhập hàng"]
                    )

                    can_ban_val = (
                        pd.to_numeric(df_bbnt[col_can_ban], errors="coerce")
                        .fillna(0)
                        .astype(float)
                    )
                    res_df["Đối chiếu KL TC & KL BBNT"] = (
                        res_df["KL theo bc trạm cân"] == can_ban_val
                    )

                    gross_ct_val = (
                        pd.to_numeric(df_bbnt[col_gross_ct], errors="coerce")
                        .replace(0, pd.NA)
                        .fillna(1)
                        .astype(float)
                    )
                    res_df["Chênh lệch (kg)"] = (
                        res_df["KL theo bc trạm cân"] - gross_ct_val
                    )
                    res_df["Chênh lệch dung sai"] = (
                        res_df["Chênh lệch (kg)"] / gross_ct_val
                    )

                    lech_pt_val = (
                        pd.to_numeric(
                            df_bbnt[col_lech_pt_bbnt], errors="coerce"
                        )
                        .fillna(0)
                        .astype(float)
                    )
                    res_df["Đối chiếu dung sai"] = (
                        res_df["Chênh lệch dung sai"].round(6)
                        == lech_pt_val.round(6)
                    )
                    res_df["So với quy định ( >+-0.3%)"] = res_df[
                        "Chênh lệch dung sai"
                    ].apply(
                        lambda x: "VUOT DUNG SAI" if abs(x) > 0.003 else "OK"
                    )

                    res_df["KL Nhập kho 035"] = ma_hoa.map(
                        lambda x: opm_035_po.get(x, 0)
                        - abs(opm_035_rtv.get(x, 0))
                    )
                    res_df["Chênh lệch "] = (
                        res_df["Khối lượng net theo PL"]
                        - res_df["KL Nhập kho 035"]
                    )

                    res_df["KL Nhập kho 031"] = ma_hoa.map(
                        lambda x: opm_031.get(x, 0)
                    )
                    res_df["Chênh lệch  "] = res_df.apply(
                        lambda r: 0
                        if r["KL Nhập kho 031"] == 0
                        else r["Khối lượng net theo PL"]
                        - r["KL Nhập kho 031"],
                        axis=1,
                    )

                    res_df["KL Nhập kho 03G-3H"] = ma_hoa.map(
                        lambda x: opm_03g.get(x, 0)
                    )
                    res_df["Chênh lệch   "] = (
                        res_df["Khối lượng net theo PL"]
                        - res_df["KL Nhập kho 03G-3H"]
                    )

                    # 7. XUẤT EXCEL & TÔ MÀU THEO CHUẨN KSNB
                    output = io.BytesIO()
                    with pd.ExcelWriter(
                        output, engine="openpyxl"
                    ) as writer:
                        res_df.to_excel(
                            writer, index=False, sheet_name="BBNT_Doi_Chieu"
                        )
                        ws = writer.sheets["BBNT_Doi_Chieu"]

                        yellow_fill = PatternFill(
                            start_color="FFFFFF00",
                            end_color="FFFFFF00",
                            fill_type="solid",
                        )
                        red_fill = PatternFill(
                            start_color="FFFFC7CE",
                            end_color="FFFFC7CE",
                            fill_type="solid",
                        )
                        red_font = Font(color="9C0006", bold=True)
                        border_thin = Border(
                            left=Side(style="thin", color="D9D9D9"),
                            right=Side(style="thin", color="D9D9D9"),
                            top=Side(style="thin", color="D9D9D9"),
                            bottom=Side(style="thin", color="D9D9D9"),
                        )

                        # Tô màu vàng các cột kiểm soát thêm mới
                        for c in range(14, len(res_df.columns) + 1):
                            cell = ws.cell(row=1, column=c)
                            cell.fill = yellow_fill
                            cell.font = Font(bold=True)

                        col_names = list(res_df.columns)
                        for r_idx in range(2, len(res_df) + 2):
                            for c_idx in range(1, len(col_names) + 1):
                                cell = ws.cell(row=r_idx, column=c_idx)
                                c_name = col_names[c_idx - 1]
                                val = cell.value

                                cell.border = border_thin
                                is_error = False

                                if (
                                    c_name
                                    in [
                                        "Đối chiếu số phiếu cân",
                                        "Đối chiếu KL TC & KL BCBV",
                                        "Đối chiếu KL TC & KL BBNT",
                                        "Đối chiếu dung sai",
                                    ]
                                    and val is False
                                ):
                                    is_error = True
                                elif (
                                    c_name == "So với quy định ( >+-0.3%)"
                                    and str(val).strip() == "VUOT DUNG SAI"
                                ):
                                    is_error = True
                                elif (
                                    c_name in ["Đối chiếu", "Chênh lệch "]
                                    and isinstance(val, (int, float))
                                    and val != 0
                                ):
                                    is_error = True

                                if is_error:
                                    cell.fill = red_fill
                                    cell.font = red_font

                    st.success("✅ Đã đối soát thành công và khớp dữ liệu!")
                    st.dataframe(res_df, use_container_width=True)

                    st.download_button(
                        "📥 Tải File Kết Quả Nghiệm Thu (Đã Tô Đỏ Lỗi)",
                        data=output.getvalue(),
                        file_name="BBNT_Nghiem_Thu_Da_Doi_Chieu.xlsx",
                        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    )

                except Exception as e:
                    st.error(f"Lỗi xử lý chi tiết: {e}")
    else:
        st.info("👉 Hãy nạp tối thiểu File BBNT để bắt đầu đối chiếu.")


doi_chieu_hrc_tung_file()
