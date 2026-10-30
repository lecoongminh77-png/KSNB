import sys
import os
import io
import time
import pandas as pd
import streamlit as st
import openpyxl
from openpyxl.styles import PatternFill, Font, Border, Side
import traceback

# -------------------------------------------------------------
# CẤU HÌNH GIAO DIỆN CHUẨN APPLE & THANH TIỆN ÍCH CỐ ĐỊNH
# -------------------------------------------------------------
st.set_page_config(
    page_title='Đối Chiếu Nghiệm Thu HRC - KSNB', page_icon='🛡️', layout='wide'
)

st.markdown(
    """
<style>
    [data-testid="stSidebar"] { display: none !important; }
    [data-testid="collapsedControl"] { display: none !important; }
    .stApp {
        background-color: #f5f5f7;
        font-family: -apple-system, BlinkMacSystemFont, "SF Pro Display", sans-serif;
        color: #1d1d1f;
    }
    .apple-header-box {
        display: flex;
        justify-content: space-between;
        align-items: flex-end;
        padding: 20px 0 10px 0;
        margin-bottom: 5px;
    }
    .apple-main-title {
        font-size: 40px;
        font-weight: 700;
        color: #1d1d1f;
        letter-spacing: -0.025em;
        margin: 0;
    }
    div.stButton > button:first-child {
        background: #0071e3 !important;
        color: #ffffff !important;
        border-radius: 980px !important;
        border: none !important;
        padding: 10px 20px !important;
        font-weight: 600 !important;
        font-size: 14px !important;
        box-shadow: 0 4px 12px rgba(0, 113, 227, 0.2) !important;
    }
    div.stButton > button:first-child:hover {
        background: #0077ed !important;
        color: #ffffff !important;
    }
    [data-testid="stVerticalBlockBorderWrapper"] {
        background: #ffffff;
        border-radius: 16px;
        padding: 15px;
        border: 1px solid rgba(0, 0, 0, 0.08);
        box-shadow: 0 2px 12px rgba(0, 0, 0, 0.03);
    }
    .upload-title {
        font-size: 18px;
        font-weight: 700;
        color: #1d1d1f;
        margin-bottom: 16px;
        display: flex;
        align-items: center;
        gap: 8px;
    }
    .guide-box {
        background: #fbfbfd;
        border-radius: 12px;
        padding: 16px;
        margin-bottom: 16px;
        border: 1px solid rgba(0,0,0,0.04);
    }
    .guide-header {
        font-weight: 700;
        color: #0071e3;
        font-size: 14px;
        margin-bottom: 8px;
        display: flex;
        align-items: center;
        gap: 6px;
    }
    .guide-text {
        font-size: 13px;
        color: #515154;
        line-height: 1.6;
        margin: 0;
    }
</style>
""",
    unsafe_allow_html=True,
)

col_home, _ = st.columns([1.5, 8.5])
with col_home:
    if st.button('🏠 Trang chủ', use_container_width=True):
        st.switch_page('app_ksnb.py')

st.markdown("<hr style='border: none; border-top: 1px solid #e2e8f0; margin: 10px 0 20px 0;'>", unsafe_allow_html=True)
st.markdown(
    """
<div class="apple-header-box">
    <div>
        <h1 class="apple-main-title">📑 Đối Chiếu Nghiệm Thu HRC</h1>
    </div>
</div>
""", unsafe_allow_html=True)
st.markdown("<hr style='border: none; border-top: 1px solid #e2e8f0; margin: 15px 0 25px 0;'>", unsafe_allow_html=True)

# ==============================================================================
# HÀM HỖ TRỢ
# ==============================================================================
def to_float(v):
    if v is None: return 0.0
    v_str = str(v).strip().upper()
    if v_str in ['', 'NONE', '#N/A', 'NAN']: return 0.0
    try: return float(v)
    except: return 0.0

# ==============================================================================
# CHỨC NĂNG LÕI: ĐỌC DỮ LIỆU PANDAS & BƠM TRỰC TIẾP VÀO FILE GỐC (BẢO LƯU 100% TEMPLATE)
# ==============================================================================
def doi_chieu_hrc_tung_file():
    st.caption('Kiểm soát nội bộ • Đọc số liệu chuẩn Pandas, giữ nguyên biểu mẫu gốc và kiểm toán dòng Tổng Cộng')
    
    st.markdown("### 📥 Nạp Dữ Liệu Nguồn")
    
    col_upload1, col_upload2 = st.columns(2, gap="large")
    
    with col_upload1:
        with st.container(border=True):
            st.markdown("""
<div class="upload-title">1️⃣ Dữ liệu Nghiệm thu & Trạm cân</div>
<div class="guide-box">
<div class="guide-header">📌 1. Biên bản nghiệm thu (BBNT)</div>
</div>
""", unsafe_allow_html=True)
            f_bbnt = st.file_uploader('1. File Biên bản nghiệm thu (BBNT)', type=['xlsx', 'xls'])
            
            selected_sheets = []
            if f_bbnt is not None:
                try:
                    xls_bbnt = pd.ExcelFile(f_bbnt)
                    sheet_names = xls_bbnt.sheet_names
                    selected_sheets = st.multiselect(
                        '📌 Chọn các Lô / Sheet BBNT cần đối chiếu:',
                        options=sheet_names,
                        default=[sheet_names[0]] if sheet_names else []
                    )
                except Exception as e:
                    st.error(f"Lỗi đọc danh sách sheet BBNT: {e}")
                    
            st.markdown("""
<div class="guide-box">
<div class="guide-header">📌 2. Packing List nhà cung cấp & 3. Trạm Cân</div>
</div>
""", unsafe_allow_html=True)
            f_pl = st.file_uploader('2. File Packing List nhà cung cấp', type=['xlsx', 'xls'])
            f_tc = st.file_uploader('3. File Dữ liệu Trạm cân', type=['xlsx', 'xls'])

    with col_upload2:
        with st.container(border=True):
            st.markdown("""
<div class="upload-title">2️⃣ Dữ liệu Bảo vệ & ERP</div>
<div class="guide-box">
<div class="guide-header">📌 4. Dữ liệu Nhập Bảo vệ</div>
</div>
<div class="guide-box">
<div class="guide-header">📌 5. Báo cáo ERP OPM 105</div>
</div>
""", unsafe_allow_html=True)
            f_bv = st.file_uploader('4. File Dữ liệu Nhập Bảo vệ', type=['xlsx', 'xls'])
            f_opm = st.file_uploader('5. File Báo cáo ERP OPM 105', type=['xlsx', 'xls'])

    if f_bbnt is not None:
        if len(selected_sheets) == 0:
            st.warning("⚠️ Vui lòng chọn ít nhất 1 sheet BBNT để bắt đầu đối chiếu.")
        else:
            if st.button('🚀 Bắt đầu đối chiếu tự động', use_container_width=True):
                my_bar = st.progress(0, text="🔄 Bắt đầu nạp hệ thống...")
                
                try:
                    # =========================================================
                    # BƯỚC 1: XÂY DỰNG KHO DỮ LIỆU TỪ CÁC FILE BỔ TRỢ (CHẠY 1 LẦN)
                    # =========================================================
                    my_bar.progress(10, text="📦 Đang xử lý Packing List...")
                    dict_pl_master = {}
                    if f_pl is not None:
                        xls_pl = pd.ExcelFile(f_pl)
                        pl_sheets = xls_pl.sheet_names
                        for s in selected_sheets:
                            if s in pl_sheets:
                                try:
                                    df_pl = pd.read_excel(f_pl, sheet_name=s, skiprows=6)
                                    if len(df_pl.columns) > 11:
                                        col_pl_ma = df_pl.columns[7]
                                        col_pl_net = df_pl.columns[11]
                                        df_pl_clean = df_pl.dropna(subset=[col_pl_ma, col_pl_net]).copy()
                                        pl_key = df_pl_clean[col_pl_ma].astype(str).str.replace(r'\.0$', '', regex=True).str.strip().str.upper() + 'A'
                                        pl_net_val = pd.to_numeric(df_pl_clean[col_pl_net], errors='coerce').fillna(0)
                                        dict_pl_master[s] = dict(zip(pl_key, pl_net_val))
                                except Exception: pass

                    my_bar.progress(25, text="⚖️ Đang xử lý Dữ liệu Trạm cân & Bảo vệ...")
                    dict_tc_kl, dict_tc_phanloai = {}, {}
                    if f_tc is not None:
                        try:
                            df_tc = pd.read_excel(f_tc)
                            tc_p = df_tc['PHIẾU CÂN'].astype(str).str.replace(r'\.0$', '', regex=True).str.strip()
                            tc_kl = pd.to_numeric(df_tc['KHỐI LƯỢNG PHIẾU CÂN'], errors='coerce').fillna(0)
                            tc_pl = df_tc['PHÂN LOẠI'].fillna('Nhập hàng').astype(str).str.strip() if 'PHÂN LOẠI' in df_tc.columns else 'Nhập hàng'
                            dict_tc_kl = dict(zip(tc_p, tc_kl))
                            dict_tc_phanloai = dict(zip(tc_p, tc_pl)) if not isinstance(tc_pl, str) else {k: tc_pl for k in tc_p}
                        except Exception: pass

                    dict_bv_phieu, dict_bv_kl = {}, {}
                    if f_bv is not None:
                        try:
                            df_bv = pd.read_excel(f_bv)
                            for _, r in df_bv.iterrows():
                                p_can = str(r.get('MÃ PHIẾU CÂN', '')).replace('.0', '').strip()
                                txt = f"{r.get('MÃ HÀNG HÓA', '')} {r.get('TÊN HÀNG HÓA', '')}".upper()
                                dict_bv_phieu[txt] = p_can 
                                dict_bv_kl[txt] = r.get('KHỐI LƯỢNG PHIẾU CÂN', 0)
                        except Exception: pass

                    my_bar.progress(40, text="🏭 Đang xử lý Báo cáo ERP OPM...")
                    opm_035_po, opm_035_rtv, opm_031, opm_03g = {}, {}, {}, {}
                    if f_opm is not None:
                        try:
                            df_opm = pd.read_excel(f_opm, skiprows=8)
                            df_opm_clean = df_opm.dropna(subset=[df_opm.columns[0], df_opm.columns[1], df_opm.columns[10]]).copy()
                            df_opm_clean['Mã Lô'] = df_opm_clean[df_opm.columns[0]].astype(str).str.strip().str.upper()
                            df_opm_clean['Khối Lượng'] = pd.to_numeric(df_opm_clean[df_opm.columns[10]], errors='coerce').fillna(0)
                            df_opm_clean['Mã Kho'] = df_opm_clean[df_opm.columns[35]].astype(str).str.strip()
                            
                            opm_035_po = df_opm_clean[(df_opm_clean['Mã Kho'] == '035') & (df_opm_clean[df_opm.columns[1]] == 'PO Receipt')].groupby('Mã Lô')['Khối Lượng'].sum().to_dict()
                            opm_035_rtv = df_opm_clean[(df_opm_clean['Mã Kho'] == '035') & (df_opm_clean[df_opm.columns[1]] == 'Return to Vendor')].groupby('Mã Lô')['Khối Lượng'].sum().to_dict()
                            opm_031 = df_opm_clean[(df_opm_clean['Mã Kho'].isin(['03G', '03H'])) & (df_opm_clean[df_opm.columns[1]] == 'Intransit Receipt')].groupby('Mã Lô')['Khối Lượng'].sum().to_dict()
                            opm_03g = df_opm_clean[(df_opm_clean['Mã Kho'] == '03G') & (df_opm_clean[df_opm.columns[1]] == 'PO Receipt')].groupby('Mã Lô')['Khối Lượng'].sum().to_dict()
                        except Exception: pass

                    # =========================================================
                    # BƯỚC 2: MỞ FILE GỐC (ĐỂ GIỮ TEMPLATE) VÀ BƠM DỮ LIỆU
                    # =========================================================
                    my_bar.progress(50, text="🧠 Đang phân tích kép biểu mẫu BBNT...")
                    
                    # Dùng openpyxl để giữ nguyên 100% template, merge cell, chú thích...
                    wb = openpyxl.load_workbook(f_bbnt, data_only=False) 
                    
                    yellow_fill = PatternFill(start_color='FFFFFF00', end_color='FFFFFF00', fill_type='solid')
                    red_fill = PatternFill(start_color='FFFFC7CE', end_color='FFFFC7CE', fill_type='solid')
                    red_font = Font(color='9C0006', bold=True)
                    border_thin = Border(left=Side(style='thin', color='D9D9D9'), right=Side(style='thin', color='D9D9D9'), top=Side(style='thin', color='D9D9D9'), bottom=Side(style='thin', color='D9D9D9'))

                    new_cols = [
                        'NHẬP/XUẤT', 'Mã hóa', 'Khối lượng net theo PL', 'Đối chiếu', 
                        'KL theo bc nhập hàng', 'Số phiếu cân', 'Đối chiếu số phiếu cân', 
                        'KL theo bc trạm cân', 'Đối chiếu KL TC & KL BCBV', 'Đối chiếu KL TC & KL BBNT', 
                        'Chênh lệch (kg)', 'Chênh lệch dung sai', 'Đối chiếu dung sai', 
                        'So với quy định ( >+-0.3%)', 'KL Nhập kho 035', 'Chênh lệch ', 
                        'KL Nhập kho 031', 'Chênh lệch  ', 'KL Nhập kho 03G-3H', 'Chênh lệch   '
                    ]

                    total_sheets = len(selected_sheets)
                    processed_count = 0
                    
                    for sheet_name in list(wb.sheetnames):
                        if sheet_name not in selected_sheets:
                            wb.remove(wb[sheet_name]) # Xóa sheet thừa cho gọn
                            continue
                            
                        processed_count += 1
                        current_progress = 50 + int((processed_count / total_sheets) * 40)
                        my_bar.progress(min(current_progress, 95), text=f"🔍 Đang đối chiếu Sheet: {sheet_name}...")

                        # --- ĐỌC SỐ LIỆU BẰNG PANDAS (Dùng nguyên logic cũ của bạn) ---
                        df_raw = pd.read_excel(f_bbnt, sheet_name=sheet_name, skiprows=7)
                        
                        # Tìm dòng Tổng cộng
                        tong_idx = -1
                        for idx, row_vals in df_raw.iterrows():
                            if 'tổng cộng' in ' '.join([str(x).lower() for x in row_vals if pd.notna(x)]):
                                tong_idx = idx
                                break
                                
                        if tong_idx != -1:
                            df_bbnt = df_raw.iloc[:tong_idx].copy()
                        else:
                            df_bbnt = df_raw.copy()

                        # Loại bỏ các dòng trống
                        df_bbnt = df_bbnt.dropna(subset=[df_bbnt.columns[3]]).copy()

                        # Logic cốt lõi (Không thay đổi)
                        base_cols = df_bbnt.columns[:13].tolist()
                        res_df = df_bbnt[base_cols].copy()
                        
                        col_ma_cuon, col_phieu_can = df_bbnt.columns[3], df_bbnt.columns[4]
                        col_net_ct, col_gross_ct = df_bbnt.columns[7], df_bbnt.columns[8]
                        col_can_ban, col_lech_pt_bbnt = df_bbnt.columns[9], df_bbnt.columns[11]

                        ma_cuon_clean = df_bbnt[col_ma_cuon].astype(str).str.replace(r'\.0$', '', regex=True).str.strip().str.upper()
                        phieu_can_clean = df_bbnt[col_phieu_can].astype(str).str.replace(r'\.0$', '', regex=True).str.strip()
                        ma_hoa = ma_cuon_clean + 'A'
                        
                        dict_pl = dict_pl_master.get(sheet_name, {})

                        res_df['NHẬP/XUẤT'] = phieu_can_clean.map(lambda x: dict_tc_phanloai.get(x, 'Nhập hàng'))
                        res_df['Mã hóa'] = ma_hoa
                        res_df['Khối lượng net theo PL'] = ma_hoa.map(dict_pl).fillna(0)

                        net_ct_val = pd.to_numeric(df_bbnt[col_net_ct], errors='coerce').fillna(0)
                        res_df['Đối chiếu'] = (net_ct_val - res_df['Khối lượng net theo PL']).astype(float)

                        res_df['KL theo bc nhập hàng'] = ma_cuon_clean.map(dict_bv_kl).fillna(0)
                        res_df['Số phiếu cân'] = ma_cuon_clean.map(dict_bv_phieu).fillna('')
                        res_df['Đối chiếu số phiếu cân'] = phieu_can_clean == res_df['Số phiếu cân'].astype(str)

                        res_df['KL theo bc trạm cân'] = phieu_can_clean.map(dict_tc_kl).fillna(0).astype(float)
                        res_df['Đối chiếu KL TC & KL BCBV'] = res_df['KL theo bc trạm cân'] == pd.to_numeric(res_df['KL theo bc nhập hàng'], errors='coerce')

                        can_ban_val = pd.to_numeric(df_bbnt[col_can_ban], errors='coerce').fillna(0).astype(float)
                        res_df['Đối chiếu KL TC & KL BBNT'] = res_df['KL theo bc trạm cân'] == can_ban_val

                        gross_ct_val = pd.to_numeric(df_bbnt[col_gross_ct], errors='coerce').replace(0, pd.NA).fillna(1).astype(float)
                        res_df['Chênh lệch (kg)'] = res_df['KL theo bc trạm cân'] - gross_ct_val
                        res_df['Chênh lệch dung sai'] = res_df['Chênh lệch (kg)'] / gross_ct_val

                        lech_pt_val = pd.to_numeric(df_bbnt[col_lech_pt_bbnt], errors='coerce').fillna(0)
                        res_df['Đối chiếu dung sai'] = res_df['Chênh lệch dung sai'].round(6) == lech_pt_val.round(6)
                        res_df['So với quy định ( >+-0.3%)'] = res_df['Chênh lệch dung sai'].apply(lambda x: 'VUOT DUNG SAI' if abs(x) > 0.003 else 'OK')

                        res_df['KL Nhập kho 035'] = ma_hoa.map(lambda x: opm_035_po.get(x, 0) - abs(opm_035_rtv.get(x, 0)))
                        res_df['Chênh lệch '] = res_df['Khối lượng net theo PL'] - res_df['KL Nhập kho 035']

                        res_df['KL Nhập kho 031'] = ma_hoa.map(lambda x: opm_031.get(x, 0))
                        res_df['Chênh lệch  '] = res_df.apply(lambda r: 0 if r['KL Nhập kho 031'] == 0 else r['Khối lượng net theo PL'] - r['KL Nhập kho 031'], axis=1)

                        res_df['KL Nhập kho 03G-3H'] = ma_hoa.map(lambda x: opm_03g.get(x, 0))
                        res_df['Chênh lệch   '] = res_df['Khối lượng net theo PL'] - res_df['KL Nhập kho 03G-3H']

                        # --- BƠM KẾT QUẢ VÀO FILE GỐC (OPENPYXL) ---
                        ws = wb[sheet_name]
                        
                        # Tô vàng Header Cột đối chiếu (Dòng 8)
                        for i, col_name in enumerate(new_cols):
                            cell = ws.cell(row=8, column=14+i)
                            cell.value = col_name
                            cell.fill = yellow_fill
                            cell.font = Font(bold=True)
                            cell.border = border_thin

                        # Ghi Data vào đúng hàng của Excel
                        for idx in res_df.index:
                            excel_row = idx + 9 # Vì Pandas skiprows=7, Index 0 tương đương dòng 9 Excel
                            
                            for c_idx, col_name in enumerate(new_cols):
                                val = res_df.loc[idx, col_name]
                                cell = ws.cell(row=excel_row, column=14+c_idx)
                                cell.value = val
                                cell.border = border_thin
                                
                                # Check lỗi tô đỏ
                                is_error = False
                                if col_name in ['Đối chiếu số phiếu cân', 'Đối chiếu KL TC & KL BCBV', 'Đối chiếu KL TC & KL BBNT', 'Đối chiếu dung sai'] and val is False:
                                    is_error = True
                                elif col_name == 'So với quy định ( >+-0.3%)' and str(val).strip() == 'VUOT DUNG SAI':
                                    is_error = True
                                elif col_name in ['Đối chiếu', 'Chênh lệch ', 'Chênh lệch  ', 'Chênh lệch   '] and isinstance(val, (int, float)) and val != 0:
                                    is_error = True
                                    
                                if is_error:
                                    cell.fill = red_fill
                                    cell.font = red_font

                        # --- KIỂM TOÁN LẠI DÒNG TỔNG CỘNG ---
                        if tong_idx != -1:
                            tong_excel_row = tong_idx + 9
                            
                            # Tự tính tổng từ dữ liệu chuẩn
                            chenh_lech_kg_file = pd.to_numeric(res_df.iloc[:, 10], errors='coerce').fillna(0)
                            manual_net = net_ct_val.sum()
                            manual_gross = gross_ct_val.sum()
                            manual_canban = can_ban_val.sum()
                            manual_chenhlech = chenh_lech_kg_file.sum()
                            manual_pct = (manual_chenhlech / manual_gross) if manual_gross != 0 else 0
                            
                            # Lấy số Tổng Cộng mà file đang tính
                            file_net = to_float(df_raw.iloc[tong_idx, 7])
                            file_gross = to_float(df_raw.iloc[tong_idx, 8])
                            file_canban = to_float(df_raw.iloc[tong_idx, 9])
                            file_chenhlech = to_float(df_raw.iloc[tong_idx, 10])
                            file_pct = to_float(df_raw.iloc[tong_idx, 11])
                            
                            # Nếu Kế toán kéo SUM sai, tô đỏ nguyên dòng đó trên Excel
                            has_sum_error = False
                            if (abs(manual_net - file_net) > 0.1 or 
                                abs(manual_gross - file_gross) > 0.1 or 
                                abs(manual_canban - file_canban) > 0.1 or 
                                abs(manual_chenhlech - file_chenhlech) > 0.1 or 
                                abs(manual_pct - file_pct) > 0.0001):
                                has_sum_error = True
                                
                            if has_sum_error:
                                for c_idx in range(1, 14):
                                    c_cell = ws.cell(row=tong_excel_row, column=c_idx)
                                    c_cell.fill = red_fill
                                    c_cell.font = red_font

                    # --- LƯU FILE KẾT QUẢ ---
                    my_bar.progress(98, text="💾 Đang đóng gói File Excel...")
                    output = io.BytesIO()
                    wb.save(output)
                    
                    my_bar.progress(100, text="✅ Hoàn tất 100%!")
                    st.success(f'🎉 Đã đối chiếu hoàn tất {len(selected_sheets)} lô hàng (File gốc được bảo lưu trọn vẹn)!')
                    st.download_button(
                        '📥 Tải File Biên Bản Đối Chiếu',
                        data=output.getvalue(),
                        file_name='BBNT_Nghiem_Thu_Doi_Chieu_Tong_Hop.xlsx',
                        mime='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
                    )

                except Exception as e:
                    my_bar.empty()
                    st.error(f'🛑 Lỗi Cốt Lõi: {e}')
                    with st.expander("🔍 Chi tiết lỗi (Dành cho kỹ thuật)"):
                        st.code(traceback.format_exc())
    else:
        st.info('👉 Hãy nạp tối thiểu File BBNT để bắt đầu.')

doi_chieu_hrc_tung_file()
