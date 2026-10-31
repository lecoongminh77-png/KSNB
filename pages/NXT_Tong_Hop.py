import io
import re
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
import pandas as pd
import streamlit as st

# -------------------------------------------------------------
# 1. CẤU HÌNH GIAO DIỆN CHUẨN APPLE
# -------------------------------------------------------------
st.set_page_config(
    page_title="Đối Chiếu Nghiệm Thu HRC - KSNB",
    page_icon="🛡️",
    layout="wide",
)

st.markdown(
    """
<style>
    [data-testid="stSidebar"] { display: none !important; }
    [data-testid="collapsedControl"] { display: none !important; }
    .stApp {
        background-color: #f5f5f7;
        font-family: -apple-system, BlinkMacSystemFont, "SF Pro Display", sans-serif;
        color: #1d1d1f;
    }
    .apple-header-box {
        padding: 20px 0 10px 0;
        margin-bottom: 5px;
    }
    .apple-main-title {
        font-size: 32px;
        font-weight: 700;
        color: #1d1d1f;
        letter-spacing: -0.02em;
        margin: 0;
    }
    div.stButton > button:first-child {
        background: #0071e3 !important;
        color: #ffffff !important;
        border-radius: 980px !important;
        border: none !important;
        padding: 12px 24px !important;
        font-weight: 600 !important;
        font-size: 15px !important;
        box-shadow: 0 4px 12px rgba(0, 113, 227, 0.25) !important;
    }
    div.stButton > button:first-child:hover {
        background: #0077ed !important;
    }
    [data-testid="stVerticalBlockBorderWrapper"] {
        background: #ffffff;
        border-radius: 16px;
        padding: 18px;
        border: 1px solid rgba(0, 0, 0, 0.08);
        box-shadow: 0 2px 10px rgba(0, 0, 0, 0.03);
    }
</style>
""",
    unsafe_allow_html=True,
)

st.markdown(
    """
<div class="apple-header-box">
    <h1 class="apple-main-title">📑 Đối Chiếu Nghiệm Thu HRC - KSNB</h1>
    <p style="color: #6e6e73; font-size: 14px; margin-top: 4px;">Tự động bắt cặp sheet, scan dữ liệu 5 bên và lọc xuất đúng các sheet cần dùng</p>
</div>
<hr style='border: none; border-top: 1px solid #d2d2d7; margin: 10px 0 20px 0;'>
""",
    unsafe_allow_html=True,
)


# -------------------------------------------------------------
# 2. CÁC HÀM XỬ LÝ SCAN DỮ LIỆU
# -------------------------------------------------------------
def tim_dong_tieu_de(file_obj, tu_khoa_list, sheet_name=0, max_scan=20):
    """Tự tìm dòng header thực tế dựa vào từ khóa."""
    file_obj.seek(0)
    df_raw = pd.read_excel(
        file_obj, sheet_name=sheet_name, header=None, nrows=max_scan
    )
    for idx, row in df_raw.iterrows():
        row_text = " ".join([str(v).upper() for v in row.values if pd.notna(v)])
        if any(k.upper() in row_text for k in tu_khoa_list):
            return idx
    return 0


def xu_ly_doi_chieu_da_sheet(f_bb, f_pl, f_tc, f_bv, f_opm, danh_sach_sheets):
    # Đọc workbook BBNT (bản công thức để ghi đè và bản data_only để lấy số)
    f_bb.seek(0)
    wb = load_workbook(f_bb, data_only=False)
    f_bb.seek(0)
    wb_val = load_workbook(f_bb, data_only=True)

    # 1. SCAN DỮ LIỆU BẢO VỆ (Quét chung 1 lần)
    dict_bv_phieu = {}
    dict_bv_kl = {}
    if f_bv:
        f_bv.seek(0)
        df_bv = pd.read_excel(f_bv)
        col_pcan = next(
            (
                c
                for c in df_bv.columns
                if "MÃ PHIẾU CÂN" in str(c).upper()
                or "PHIẾU CÂN" in str(c).upper()
            ),
            None,
        )
        col_klcan = next(
            (c for c in df_bv.columns if "KHỐI LƯỢNG" in str(c).upper()), None
        )
        col_hang = next(
            (
                c
                for c in df_bv.columns
                if "MÃ HÀNG HÓA" in str(c).upper()
                or "TÊN HÀNG HÓA" in str(c).upper()
            ),
            None,
        )

        # Gom toàn bộ dữ liệu bảo vệ để tra cứu nhanh
        bv_records = []
        for _, r in df_bv.iterrows():
            txt = f"{r.get('MÃ HÀNG HÓA', '')} {r.get('TÊN HÀNG HÓA', '')}"
            p_can = str(r.get(col_pcan, "")).replace(".0", "").strip()
            kl_can = r.get(col_klcan, 0)
            bv_records.append((txt, p_can, kl_can))

    # 2. SCAN DỮ LIỆU TRẠM CÂN (Quét chung 1 lần)
    dict_tc_kl = {}
    dict_tc_pl = {}
    if f_tc:
        f_tc.seek(0)
        df_tc = pd.read_excel(f_tc)
        col_p = next(
            c for c in df_tc.columns if str(c).strip().upper() == "PHIẾU CÂN"
        )
        col_k = next(c for c in df_tc.columns if "KHỐI LƯỢNG" in str(c).upper())
        col_pl = next(
            (c for c in df_tc.columns if "PHÂN LOẠI" in str(c).upper()), None
        )

        p_series = (
            df_tc[col_p]
            .astype(str)
            .str.replace(r"\.0$", "", regex=True)
            .str.strip()
        )
        kl_series = pd.to_numeric(df_tc[col_k], errors="coerce").fillna(0)
        dict_tc_kl = dict(zip(p_series, kl_series))
        if col_pl:
            dict_tc_pl = dict(
                zip(p_series, df_tc[col_pl].fillna("Nhập hàng").astype(str))
            )

    # 3. SCAN DỮ LIỆU ERP OPM 105 (Quét chung 1 lần)
    opm_035_po, opm_035_rtv, opm_031, opm_03g_po, opm_03g_rtv = (
        {},
        {},
        {},
        {},
        {},
    )
    if f_opm:
        f_opm.seek(0)
        idx_opm = tim_dong_tieu_de(f_opm, ["LOẠI GIAO DỊCH", "MÃ LÔ"])
        f_opm.seek(0)
        df_opm = pd.read_excel(f_opm, skiprows=idx_opm)
        if len(df_opm) > 0 and "[" in str(df_opm.iloc[0, 1]):
            df_opm = df_opm.iloc[1:].copy()

        c_loai = next(
            c for c in df_opm.columns if "LOẠI GIAO DỊCH" in str(c).upper()
        )
        c_malo = next(c for c in df_opm.columns if "MÃ LÔ" in str(c).upper())
        c_kl = next(
            c
            for c in df_opm.columns
            if "KHỐI LƯỢNG" in str(c).upper() and "GROSS" not in str(c).upper()
        )
        c_kho = next(
            c
            for c in df_opm.columns
            if "MĂ KHO" in str(c).upper() or str(c).strip().upper() == "MÃ KHO"
        )

        df_opm_clean = df_opm.dropna(subset=[c_malo, c_loai]).copy()
        raw_lo = df_opm_clean[c_malo].astype(str).str.strip().str.upper()
        df_opm_clean["Key"] = raw_lo.apply(
            lambda x: x if x.endswith("A") else x + "A"
        )
        df_opm_clean["KL_Val"] = pd.to_numeric(
            df_opm_clean[c_kl], errors="coerce"
        ).fillna(0)
        df_opm_clean["Kho_Val"] = (
            df_opm_clean[c_kho].astype(str).str.strip().str.upper()
        )

        opm_035_po = (
            df_opm_clean[
                (df_opm_clean["Kho_Val"] == "035")
                & (
                    df_opm_clean[c_loai].str.contains(
                        "PO Receipt", case=False, na=False
                    )
                )
            ]
            .groupby("Key")["KL_Val"]
            .sum()
            .to_dict()
        )
        opm_035_rtv = (
            df_opm_clean[
                (df_opm_clean["Kho_Val"] == "035")
                & (
                    df_opm_clean[c_loai].str.contains(
                        "Return to Vendor", case=False, na=False
                    )
                )
            ]
            .groupby("Key")["KL_Val"]
            .sum()
            .to_dict()
        )
        opm_031 = (
            df_opm_clean[
                (df_opm_clean["Kho_Val"].isin(["03G", "03H"]))
                & (
                    df_opm_clean[c_loai].str.contains(
                        "Intransit Receipt", case=False, na=False
                    )
                )
            ]
            .groupby("Key")["KL_Val"]
            .sum()
            .to_dict()
        )
        opm_03g_po = (
            df_opm_clean[
                (df_opm_clean["Kho_Val"] == "03G")
                & (
                    df_opm_clean[c_loai].str.contains(
                        "PO Receipt", case=False, na=False
                    )
                )
            ]
            .groupby("Key")["KL_Val"]
            .sum()
            .to_dict()
        )
        opm_03g_rtv = (
            df_opm_clean[
                (df_opm_clean["Kho_Val"] == "03G")
                & (
                    df_opm_clean[c_loai].str.contains(
                        "Return to Vendor", case=False, na=False
                    )
                )
            ]
            .groupby("Key")["KL_Val"]
            .sum()
            .to_dict()
        )

    # Đọc danh sách sheet của Packing List để bắt cặp
    f_pl.seek(0)
    xl_pl = pd.ExcelFile(f_pl)
    pl_sheets = xl_pl.sheet_names

    headers = [
        (14, "NHẬP/XUẤT"),
        (15, "Mã hóa"),
        (16, "Khối lượng net theo PL"),
        (17, "Đối chiếu "),
        (18, "KL theo bc nhập hàng "),
        (19, "Số phiếu cân "),
        (20, "Đối chiếu số phiếu cân"),
        (21, "KL theo bc trạm cân"),
        (22, ""),
        (23, "Đối chiếu KL TC & KL BCBV"),
        (24, "Đối chiếu KL TC & KL BBNT"),
        (25, "Chênh lệch (kg)"),
        (26, " Chênh lệch dung sai"),
        (27, "Đối chiếu dung sai"),
        (28, " So với quy định ( >+-0.3%)"),
        (29, "KL Nhập kho\n035"),
        (30, "Chênh lệch"),
        (31, "KL Nhập kho\n031"),
        (32, "Chênh lệch"),
        (33, "KL Nhập kho\n03G-3H"),
        (34, "Chênh lệch"),
    ]

    yellow_fill = PatternFill(
        start_color="FFFFFF00", end_color="FFFFFF00", fill_type="solid"
    )
    red_fill = PatternFill(
        start_color="FFFFC7CE", end_color="FFFFC7CE", fill_type="solid"
    )
    red_font = Font(color="9C0006", bold=True)
    border_thin = Border(
        left=Side(style="thin", color="D9D9D9"),
        right=Side(style="thin", color="D9D9D9"),
        top=Side(style="thin", color="D9D9D9"),
        bottom=Side(style="thin", color="D9D9D9"),
    )

    # 4. CHẠY TỪNG SHEET ĐƯỢC CHỌN
    for s_name in danh_sach_sheets:
        ws = wb[s_name]
        ws_val = wb_val[s_name]

        # Scan Packing list của riêng sheet này
        dict_pl = {}
        matched_pl_sheet = next(
            (p for p in pl_sheets if p.strip().upper() == s_name.strip().upper()),
            None,
        )
        if matched_pl_sheet:
            idx_pl = tim_dong_tieu_de(
                f_pl,
                ["MÃ SỐ NSX", "MÃ SỐ HSPM", "KL(NET)"],
                sheet_name=matched_pl_sheet,
            )
            f_pl.seek(0)
            df_pl = pd.read_excel(
                f_pl, sheet_name=matched_pl_sheet, skiprows=idx_pl
            )
            col_nsx = next(
                c
                for c in df_pl.columns
                if "MÃ SỐ NSX" in str(c).upper() or "MÃ SỐ" in str(c).upper()
            )
            col_net = next(c for c in df_pl.columns if "NET" in str(c).upper())
            df_pl_clean = df_pl.dropna(subset=[col_nsx, col_net])
            keys = (
                df_pl_clean[col_nsx]
                .astype(str)
                .str.replace(r"\.0$", "", regex=True)
                .str.strip()
                + "A"
            )
            nets = pd.to_numeric(df_pl_clean[col_net], errors="coerce").fillna(0)
            dict_pl = (
                pd.DataFrame({"Key": keys, "Net": nets})
                .groupby("Key")["Net"]
                .sum()
                .to_dict()
            )

        # Tìm dòng header
        row_header = 8
        for r in range(1, 15):
            if (
                ws.cell(row=r, column=4).value
                and "CUỘN" in str(ws.cell(row=r, column=4).value).upper()
            ):
                row_header = r
                break

        # Đọc dữ liệu dòng
        data_rows = []
        for r in range(row_header + 1, ws.max_row + 1):
            c_val = ws_val.cell(row=r, column=4).value
            stt_val = str(ws_val.cell(row=r, column=1).value or "").strip()
            if (
                c_val is not None
                and str(c_val).strip() != ""
                and "TỔNG" not in stt_val.upper()
                and "KẾT LUẬN" not in stt_val.upper()
            ):
                data_rows.append((r, str(c_val).replace(".0", "").strip()))

        # Ghi header
        for col_idx, h_text in headers:
            cell = ws.cell(row=row_header, column=col_idx)
            cell.value = h_text
            cell.fill = yellow_fill
            cell.font = Font(bold=True)
            cell.alignment = Alignment(
                wrap_text=True, horizontal="center", vertical="center"
            )

        # Điền dữ liệu đối chiếu
        for r, m in data_rows:
            m_hoa = m + "A"
            p_can_bbnt = str(ws_val.cell(row=r, column=5).value or "").strip()
            net_ct = float(ws_val.cell(row=r, column=8).value or 0)
            gross_ct = float(ws_val.cell(row=r, column=9).value or 1)
            can_ban = float(ws_val.cell(row=r, column=10).value or 0)
            dsai_bbnt = round(float(ws_val.cell(row=r, column=12).value or 0), 6)

            # Match bảo vệ
            p_can_bv, kl_bv = "", 0.0
            if f_bv:
                for txt, p, kl in bv_records:
                    if m in txt:
                        p_can_bv = p
                        kl_bv = float(kl or 0)
                        break

            nhap_xuat = dict_tc_pl.get(p_can_bbnt, "Nhập hàng")
            kl_net_pl = float(dict_pl.get(m_hoa, 0))
            doi_chieu_pl = net_ct - kl_net_pl
            chk_pcan = p_can_bbnt == p_can_bv if p_can_bv else False

            p_tra_tc = p_can_bv if p_can_bv else p_can_bbnt
            kl_tc = float(dict_tc_kl.get(p_tra_tc, 0))

            chk_tc_bv = kl_tc == kl_bv if kl_bv > 0 else False
            chk_tc_bbnt = kl_tc == can_ban if can_ban > 0 else False
            chenh_lech_kg = kl_tc - gross_ct
            chenh_lech_ds = chenh_lech_kg / gross_ct if gross_ct != 0 else 0
            chk_dsai = round(chenh_lech_ds, 6) == dsai_bbnt
            so_quy_dinh = "VUOT DUNG SAI" if abs(chenh_lech_ds) > 0.003 else "OK"

            kl_035 = float(
                opm_035_po.get(m_hoa, 0) - abs(opm_035_rtv.get(m_hoa, 0))
            )
            chenh_035 = kl_net_pl - kl_035
            kl_031 = float(opm_031.get(m_hoa, 0))
            chenh_031 = 0 if kl_031 == 0 else kl_net_pl - kl_031
            kl_03g = float(
                opm_03g_po.get(m_hoa, 0) - abs(opm_03g_rtv.get(m_hoa, 0))
            )
            chenh_03g = kl_net_pl - kl_03g

            row_values = {
                14: nhap_xuat,
                15: m_hoa,
                16: kl_net_pl,
                17: doi_chieu_pl,
                18: kl_bv,
                19: p_can_bv,
                20: chk_pcan,
                21: kl_tc,
                22: p_can_bv + "A" if p_can_bv else "",
                23: chk_tc_bv,
                24: chk_tc_bbnt,
                25: chenh_lech_kg,
                26: chenh_lech_ds,
                27: chk_dsai,
                28: so_quy_dinh,
                29: kl_035,
                30: chenh_035,
                31: kl_031,
                32: chenh_031,
                33: kl_03g,
                34: chenh_03g,
            }

            for col_idx, val in row_values.items():
                c = ws.cell(row=r, column=col_idx)
                c.value = val
                c.border = border_thin

                is_err = False
                if col_idx in [20, 23, 24, 27] and val is False:
                    is_err = True
                elif col_idx == 28 and val == "VUOT DUNG SAI":
                    is_err = True
                elif (
                    col_idx in [17, 30]
                    and isinstance(val, (int, float))
                    and val != 0
                ):
                    is_err = True

                if is_err:
                    c.fill = red_fill
                    c.font = red_font

    # 5. XÓA SẠCH TOÀN BỘ CÁC SHEET KHÔNG ĐƯỢC CHỌN
    for s_name in wb.sheetnames:
        if s_name not in danh_sach_sheets:
            wb.remove(wb[s_name])

    out_io = io.BytesIO()
    wb.save(out_io)
    return out_io.getvalue()


# -------------------------------------------------------------
# 3. GIAO DIỆN STREAMLIT
# -------------------------------------------------------------
col1, col2 = st.columns(2, gap="large")

with col1:
    with st.container(border=True):
        st.subheader("1️⃣ Biên bản & Trạm cân")
        f_bb = st.file_uploader(
            "1. File Biên bản nghiệm thu (BBNT)",
            type=["xlsx", "xls"],
            key="u_bb",
        )
        f_pl = st.file_uploader(
            "2. File Packing List NCC", type=["xlsx", "xls"], key="u_pl"
        )
        f_tc = st.file_uploader(
            "3. File Dữ liệu Trạm cân", type=["xlsx", "xls"], key="u_tc"
        )

with col2:
    with st.container(border=True):
        st.subheader("2️⃣ Bảo vệ & ERP OPM")
        f_bv = st.file_uploader(
            "4. File Nhập Bảo vệ", type=["xlsx", "xls"], key="u_bv"
        )
        f_opm = st.file_uploader(
            "5. File ERP OPM 105", type=["xlsx", "xls"], key="u_opm"
        )

# XỬ LÝ CHỌN SHEET
if f_bb is not None:
    # Lấy danh sách sheets trong file BBNT
    f_bb.seek(0)
    xl_bb_temp = pd.ExcelFile(f_bb)
    available_sheets = xl_bb_temp.sheet_names

    st.markdown("### 🎯 Chọn Sheet Cần Đối Chiếu")
    sheets_selected = st.multiselect(
        "Chọn một hoặc nhiều sheet (những sheet không chọn sẽ tự động bị loại bỏ khỏi file kết quả):",
        options=available_sheets,
        default=[available_sheets[0]] if available_sheets else [],
    )

    if not sheets_selected:
        st.warning("⚠️ Vui lòng chọn ít nhất 1 sheet để tiếp tục.")
    else:
        st.write("")
        if st.button("🚀 Bắt đầu đối chiếu tự động", use_container_width=True):
            with st.spinner(
                f"Đang xử lý {len(sheets_selected)} sheet đã chọn và lọc bỏ các sheet thừa..."
            ):
                try:
                    excel_bytes = xu_ly_doi_chieu_da_sheet(
                        f_bb, f_pl, f_tc, f_bv, f_opm, sheets_selected
                    )
                    st.success(
                        f"✅ Đã đối soát thành công {len(sheets_selected)} sheet!"
                    )

                    st.download_button(
                        label=f"📥 TẢI FILE EXCEL KẾT QUẢ ({len(sheets_selected)} SHEET)",
                        data=excel_bytes,
                        file_name="BBNT_Doi_Chieu_Ket_Qua.xlsx",
                        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                        use_container_width=True,
                    )
                except Exception as e:
                    st.error(f"❌ Có lỗi trong quá trình xử lý: {e}")
else:
    st.info("👉 Hãy nạp file Biên bản nghiệm thu (BBNT) để chọn sheet cần chạy.")
